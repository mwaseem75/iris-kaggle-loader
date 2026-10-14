"""
Read a Kaggle file into a pyarrow Table: CSV/TSV, Parquet, JSON/JSONL and Excel.

CSV has no types, so every value is read as text and each column is typed from all
of its values afterwards: a column with "00123" in it stays text, "NA" and similar
markers in an otherwise numeric column become NULL.
"""
import csv
import datetime
import io
import json
import os
import re

import pyarrow as pa

from .loader import LoadError

FORMATS = {".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".parquet": "parquet", ".pq": "parquet",
           ".json": "json", ".jsonl": "jsonl", ".ndjson": "jsonl", ".xlsx": "excel", ".xlsm": "excel"}
ENCODINGS = ("utf-8", "cp1252", "latin-1")
NULL_TOKENS = {"na", "n/a", "nan", "null", "none"}
INT64_MIN, INT64_MAX = -2 ** 63, 2 ** 63 - 1

INT = re.compile(r"[+-]?(0|[1-9][0-9]*)")
FLOAT = re.compile(r"[+-]?((0|[1-9][0-9]*)(\.[0-9]*)?|\.[0-9]+)([eE][+-]?[0-9]+)?")   # no leading zeros
DATE = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}")
TIMESTAMP = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}[ T][0-9]{2}:[0-9]{2}(:[0-9]{2}(\.[0-9]{1,6})?)?(Z| ?UTC|[+-][0-9]{2}:?[0-9]{2})?")


def file_format(name):
    """'csv', 'tsv', 'parquet', 'json', 'jsonl', 'excel' or None for files that cannot be imported."""
    return FORMATS.get(os.path.splitext(str(name).lower())[1])


def read(source, fmt, limit=None, sheet=None, partial=False):
    """(table, info) for a file path or bytes. limit: at most this many rows.
    partial: the bytes are only the start of the file (a preview); the last,
    probably cut, line is dropped. info: encoding, sheets, sheet, truncated."""
    if fmt in ("csv", "tsv"):
        return _csv(_bytes(source), "\t" if fmt == "tsv" else ",", limit, partial)
    if fmt == "jsonl":
        return _jsonl(_bytes(source), limit, partial)
    if partial:
        raise LoadError(f"{fmt} files can only be read completely")
    if fmt == "parquet":
        return _parquet(source, limit)
    if fmt == "json":
        return _json(_bytes(source), limit)
    if fmt == "excel":
        return _excel(source, limit, sheet)
    raise LoadError("This file type cannot be imported (supported: CSV, TSV, Parquet, JSON, JSONL, Excel).")


def _bytes(source):
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    with open(source, "rb") as f:
        return f.read()


def decode(data):
    """(text, encoding): UTF-8 (with or without BOM), else Windows-1252, else Latin-1."""
    for encoding in ENCODINGS:
        try:
            text = data.decode("utf-8-sig" if encoding == "utf-8" else encoding)
            return text, encoding
        except UnicodeDecodeError:
            continue
    raise LoadError("The file could not be decoded as text")


# -------------------------------------------------------------------- CSV ---
def _csv(data, delimiter, limit, partial):
    text, encoding = decode(data)
    if partial and "\n" in text:
        text = text[:text.rindex("\n") + 1]            # the last line may be cut
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        raise LoadError("The file is empty")
    except csv.Error as e:
        raise LoadError(f"Line 1 cannot be read as {'TSV' if delimiter == chr(9) else 'CSV'}: {e}")
    width = len(header)
    columns = [[] for _ in header]
    rows = 0
    try:
        for record in reader:
            if not record or record == [""] * len(record) and len(record) <= 1:
                continue                                  # blank line
            if len(record) != width:
                raise LoadError(f"Line {reader.line_num} has {len(record)} fields, the header has {width}")
            for column, value in zip(columns, record):
                column.append(value if value != "" else None)
            rows += 1
            if limit and rows >= limit:
                break
    except csv.Error as e:
        raise LoadError(f"Line {reader.line_num} cannot be read: {e}")
    arrays = [infer(values) for values in columns]
    return pa.table(arrays, names=[h.strip() for h in header]), {"encoding": encoding, "truncated": partial}


def infer(values):
    """An Arrow array for a column of strings (None = empty): int64, double, bool, date,
    timestamp, or string. A column is only numeric if every value is."""
    present = [v.strip() for v in values if v is not None and v.strip() != ""]
    if not present:
        return pa.array([None] * len(values), pa.string())
    real = [v for v in present if v.lower() not in NULL_TOKENS]
    nullish = lambda v: v is None or v.strip() == "" or v.strip().lower() in NULL_TOKENS
    if real and all(INT.fullmatch(v) for v in real) and all(INT64_MIN <= int(v) <= INT64_MAX for v in real):
        return pa.array([None if nullish(v) else int(v) for v in values], pa.int64())
    if real and all(INT.fullmatch(v) or FLOAT.fullmatch(v) for v in real):
        return pa.array([None if nullish(v) else float(v) for v in values], pa.float64())
    if all(v.lower() in ("true", "false") for v in present):
        return pa.array([None if v is None or v.strip() == "" else v.strip().lower() == "true" for v in values], pa.bool_())
    if all(DATE.fullmatch(v) for v in present):
        try:
            return pa.array([None if v is None or v.strip() == "" else datetime.date.fromisoformat(v.strip()) for v in values],
                            pa.date32())
        except ValueError:
            pass                                          # e.g. 2023-02-30: keep the text
    if all(TIMESTAMP.fullmatch(v) for v in present):
        try:
            return pa.array([None if v is None or v.strip() == "" else _timestamp(v.strip()) for v in values],
                            pa.timestamp("us"))
        except ValueError:
            pass
    return pa.array(values, pa.string())


def _timestamp(text):
    text = text[:-3].rstrip() + "+00:00" if text.endswith("UTC") else text.replace("Z", "+00:00")
    t = datetime.datetime.fromisoformat(text)
    if t.tzinfo:
        t = t.astimezone(datetime.timezone.utc).replace(tzinfo=None)
    return t


# ------------------------------------------------------------ JSON, JSONL ---
def _from_records(records):
    if not records:
        raise LoadError("The file contains no records")
    if not all(isinstance(r, dict) for r in records):
        raise LoadError("Each record must be a JSON object")
    try:
        return pa.Table.from_pylist(records)
    except (pa.ArrowInvalid, pa.ArrowTypeError, pa.ArrowNotImplementedError):
        names = list(dict.fromkeys(k for r in records for k in r))
        # mixed types in a column: keep each value as JSON text
        return pa.table([pa.array([None if r.get(k) is None else json.dumps(r.get(k), ensure_ascii=False, default=str)
                                   for r in records], pa.string()) for k in names], names=names)


def _json(data, limit):
    text, encoding = decode(data)
    try:
        doc = json.loads(text)
    except ValueError as e:
        raise LoadError(f"The file is not valid JSON: {e}")
    if isinstance(doc, dict):
        lists = [v for v in doc.values() if isinstance(v, list)]
        if len(lists) != 1:
            raise LoadError("The JSON file must hold a list of records (or an object with exactly one list)")
        doc = lists[0]
    if not isinstance(doc, list):
        raise LoadError("The JSON file must hold a list of records")
    return _from_records(doc[:limit] if limit else doc), {"encoding": encoding, "truncated": False}


def _jsonl(data, limit, partial):
    text, encoding = decode(data)
    lines = text.splitlines()
    if partial and lines and not text.endswith("\n"):
        lines = lines[:-1]
    records = []
    for number, line in enumerate(lines, 1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except ValueError as e:
            raise LoadError(f"Line {number} is not valid JSON: {e}")
        if limit and len(records) >= limit:
            break
    return _from_records(records), {"encoding": encoding, "truncated": partial}


# ---------------------------------------------------------------- Parquet ---
def _parquet(source, limit):
    import pyarrow.parquet as pq
    try:
        f = pq.ParquetFile(io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source)
    except Exception as e:
        raise LoadError(f"The file cannot be read as Parquet: {e}")
    if not limit:
        return f.read(), {}
    batches, rows = [], 0
    for batch in f.iter_batches(batch_size=min(limit, 65536)):
        batches.append(batch)
        rows += batch.num_rows
        if rows >= limit:
            break
    table = pa.Table.from_batches(batches, schema=f.schema_arrow) if batches else f.schema_arrow.empty_table()
    return table.slice(0, limit), {}


# ------------------------------------------------------------------ Excel ---
def _excel(source, limit, sheet):
    try:
        import warnings
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")               # "Workbook contains no default style"
            book = load_workbook(io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source,
                                 read_only=True, data_only=True)
    except Exception as e:
        raise LoadError(f"The file cannot be read as Excel: {e}")
    try:
        sheets = book.sheetnames
        name = sheet or sheets[0]
        if name not in sheets:
            raise LoadError(f"Sheet {name!r} not found. Sheets: {', '.join(sheets)}")
        rows = book[name].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise LoadError(f"Sheet {name!r} is empty")
        names = ["" if h is None else str(h).strip() for h in header]
        while names and names[-1] == "":
            names.pop()                                   # trailing empty header cells
        columns = [[] for _ in names]
        count = 0
        for row in rows:
            if row is None or all(v is None for v in row[:len(names)]):
                continue
            for i, column in enumerate(columns):
                column.append(row[i] if i < len(row) else None)
            count += 1
            if limit and count >= limit:
                break
    finally:
        book.close()
    return pa.table([_excel_column(c) for c in columns], names=names), {"sheets": sheets, "sheet": name}


def _excel_column(values):
    """Excel cells already have types; a column keeps one type only if all its cells share it.
    Columns of text only are typed like CSV columns."""
    present = [v for v in values if v is not None]
    kinds = {type(v) for v in present}
    if not present:
        return pa.array([None] * len(values), pa.string())
    if kinds == {str}:
        return infer(values)                              # numbers saved as text are common
    if kinds == {bool}:
        return pa.array(values, pa.bool_())
    if kinds <= {int} and all(INT64_MIN <= v <= INT64_MAX for v in present):
        return pa.array(values, pa.int64())
    if kinds <= {int, float}:
        return pa.array([None if v is None else float(v) for v in values], pa.float64())
    if kinds <= {datetime.datetime}:
        return pa.array(values, pa.timestamp("us"))
    if kinds <= {datetime.date}:
        return pa.array(values, pa.date32())
    return pa.array([None if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v)) for v in values],
                    pa.string())
